from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework import status
from django.shortcuts import get_object_or_404
from django.db.models import Q
from groups.models import Group, Student
from .utils import send_telegram_message_async, get_student_telegram_ids
from django.db.models import Q

class BroadcastMessageView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        user = request.user
        group_id = request.data.get('group_id')
        branch_id = request.data.get('branch_id')
        send_to_all_branches = request.data.get('send_to_all_branches', False)
        message = request.data.get('message')

        if not message:
            return Response({"error": "Xabar matni kiritilmadi"}, status=status.HTTP_400_BAD_REQUEST)

        # Base Queryset for students
        students_qs = Student.objects.all()

        # 1. Global (Barcha filiallar) - Faqat Super Admin uchun
        if send_to_all_branches:
            if user.role != 'super_admin':
                return Response({"error": "Barcha filialarga xabar yuborish faqat Super Admin uchun mumkin"}, status=status.HTTP_403_FORBIDDEN)
            # No filtering needed for students_qs

        # 2. Guruh bo'yicha
        elif group_id:
            group = get_object_or_404(Group, id=group_id)
            # Permission check...
            has_permission = False
            if user.role == 'super_admin': has_permission = True
            elif user.role == 'admin':
                allowed_branches = [user.branch.id] if user.branch else []
                allowed_branches.extend(user.branch_accesses.values_list('branch_id', flat=True))
                if group.branch_id in allowed_branches: has_permission = True
            elif user.role == 'mentor':
                if group.mentor == user or group.additional_mentors.filter(mentor=user).exists():
                    has_permission = True
            
            if not has_permission:
                return Response({"error": "Siz ushbu guruhga xabar yuborish huquqiga ega emassiz"}, status=status.HTTP_403_FORBIDDEN)
            
            students_qs = group.students.all()

        # 3. Filial bo'yicha
        elif branch_id:
            # Permission check
            if user.role != 'super_admin':
                # Admin faqat o'ziga ruxsat berilgan filialga yubora oladi
                allowed_branches = [user.branch.id] if user.branch else []
                allowed_branches.extend(user.branch_accesses.values_list('branch_id', flat=True))
                if int(branch_id) not in allowed_branches:
                    return Response({"error": "Siz ushbu filialga xabar yubora olmaysiz"}, status=status.HTTP_403_FORBIDDEN)
            
            students_qs = students_qs.filter(branch_id=branch_id)

        # 4. Agar hech narsa berilmasa va Super Admin bo'lsa (Eski logika muvofiqligi)
        elif user.role == 'super_admin':
            pass # Barcha o'quvchilarga boradi (Global)
        
        else:
            return Response({"error": "Maqsadli guruh yoki filial ko'rsatilmadi"}, status=status.HTTP_400_BAD_REQUEST)

        # Xabarlarni yuborish uchun ID-larni yig'amiz
        target_chat_ids = set()
        for student in students_qs:
            chat_ids = get_student_telegram_ids(student)
            target_chat_ids.update(chat_ids)
        
        # Yakuniy xabar mazmuni
        final_message = message
        if group_id:
            group = Group.objects.filter(id=group_id).first()
            if group:
                final_message = f"<b>{group.name}</b>\n\n{message}"

        # Xabarlarni yuborish (Celery va Fallback Threading bilan)
        try:
            from .tasks import send_broadcast_message_task
            send_broadcast_message_task.delay(list(target_chat_ids), final_message)
        except Exception:
            # Celery/Redis ishlamayotgan bo'lsa, xabarlarni bitta fondagi thread ichida navbat bilan yuboramiz
            import threading
            import time
            from .utils import _send_message_sync

            def fallback_broadcast(ids_list, msg_text):
                # Duplikatlarni oldini olish uchun set-dan foydalanamiz
                unique_ids = list(set(ids_list))
                for cid in unique_ids:
                    try:
                        _send_message_sync(cid, msg_text)
                        time.sleep(0.05)
                    except:
                        pass
            
            threading.Thread(target=fallback_broadcast, args=(list(target_chat_ids), final_message)).start()
        
        sent_count = len(target_chat_ids)
        target_name = "Barcha o'quvchilar" if send_to_all_branches or (not group_id and not branch_id) else f"'{group.name if 'group' in locals() and group else 'Tanlangan'}' guruhi"
        return Response({
            "status": "success", 
            "detail": f"Xabar {sent_count} ta unikal foydalanuvchiga yuborilmoqda ({target_name})."
        })

class BotStatisticsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        branch_id = request.query_params.get('branch_id')
        group_id = request.query_params.get('group_id')
        
        # Faqat aktiv guruhlarda o'qiyotgan o'quvchilarni olamiz
        qs = Student.objects.filter(enrollments__is_active=True).distinct()
        
        if branch_id:
            qs = qs.filter(branch_id=branch_id)
        if group_id:
            group_obj = get_object_or_404(Group, id=group_id)
            qs = qs.filter(groups=group_obj)
            
        # O'quvchilarni sanash (telegram_id bo'sh bo'lmaganlar)
        students_bot = qs.exclude(Q(telegram_id__isnull=True) | Q(telegram_id='')).count()
        
        # Ota-onalarni sanash (parent_telegram_id bo'sh bo'lmaganlar)
        parents_bot = qs.exclude(Q(parent_telegram_id__isnull=True) | Q(parent_telegram_id='')).count()
        
        unregistered_students = qs.filter(
            (Q(telegram_id__isnull=True) | Q(telegram_id='')) &
            (Q(parent_telegram_id__isnull=True) | Q(parent_telegram_id=''))
        ).count()
        
        return Response({
            "students_bot_count": students_bot,
            "parents_bot_count": parents_bot,
            "total_bot_users": students_bot + parents_bot,
            "unregistered_students": unregistered_students
        })

class ExportUnregisteredStudentsView(APIView):
    """
    Botdan o'tmagan (chat_id si yo'q) o'quvchilar ro'yxatini Excel export qiladi.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        branch_id = request.query_params.get('branch_id')
        
        # O'quvchilar: na o'zi va na ota-onasi botdan o'tganlar (umuman ulanmaganlar)
        qs = Student.objects.filter(
            (Q(telegram_id__isnull=True) | Q(telegram_id='')) &
            (Q(parent_telegram_id__isnull=True) | Q(parent_telegram_id=''))
        ).select_related('branch').prefetch_related('groups')

        if branch_id:
            qs = qs.filter(branch_id=branch_id)

        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from django.http import HttpResponse

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Botdan o'tmaganlar"

        # Header
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="b88a0b", end_color="b88a0b", fill_type="solid") # Gold variant
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        headers = ["№", "O'quvchi ismi-familiyasi", "Guruhlari", "Telefon", "Ota-ona telefoni", "Holati"]
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        # Column widths
        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 35

        # Data
        for row_num, student in enumerate(qs, 2):
            status_text = "O'quvchi ulanmagan & Otaona ulanmagan"
            
            ws.cell(row=row_num, column=1).value = row_num - 1
            ws.cell(row=row_num, column=2).value = student.full_name
            ws.cell(row=row_num, column=3).value = ", ".join([g.name for g in student.groups.all()])
            ws.cell(row=row_num, column=4).value = student.phone or ""
            ws.cell(row=row_num, column=5).value = student.parent_phone or ""
            ws.cell(row=row_num, column=6).value = status_text
            
            # Text wrapping for groups
            ws.cell(row=row_num, column=3).alignment = Alignment(wrap_text=True)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="botdan_otmaganlar.xlsx"'
        wb.save(response)
        return response
